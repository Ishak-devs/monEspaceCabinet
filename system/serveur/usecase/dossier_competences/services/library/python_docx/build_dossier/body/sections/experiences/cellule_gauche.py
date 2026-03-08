from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

COULEUR_PRINCIPALE = RGBColor(0x1B, 0x4A, 0x8A)
POLICE = "Arial"

def cellule_gauche(table, exp):
    cell_gauche = table.cell(0, 0)
    p_gauche = cell_gauche.paragraphs[0]
    p_gauche.alignement = WD_ALIGN_PARAGRAPH.LEFT
    run = p_gauche.add_run(exp.get('Nom_Entreprise', ''))
    run.font.name = POLICE
    run.font.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = COULEUR_PRINCIPALE