from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import RGBColor, Pt

POLICE = "Arial"

def cellule_milieu(table, exp):
    cell_milieu = table.cell(0, 1)
    p_milieu = cell_milieu.paragraphs[0]
    p_milieu.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_milieu.add_run(exp.get('Durée_expérience', ''))
    run.font.name = POLICE
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x00, 0x20, 0x60)