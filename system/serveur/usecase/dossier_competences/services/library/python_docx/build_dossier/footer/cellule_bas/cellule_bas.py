from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor


def cellule_bas(tab):
    p2 = tab.cell(1, 0).paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run(
        "NAVA ENGINEERING • 9-11 AVENUE MICHELET, 93400 SAINT-OUEN-SUR-SEINE\nSimon ZANA • 06 13 53 23 81 • simon.zana@nava-eng.com")
    r.font.size, r.font.color.rgb = Pt(9), RGBColor(0, 0, 20)